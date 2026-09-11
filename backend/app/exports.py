"""Génération du classeur Excel des résultats d'une évaluation.

Le classeur est construit à partir de l'état en base : liste complète des inscrits
(y compris les absents), notes par exercice, détail des tests, historique des
ajustements et incidents d'intégrité.
"""

from __future__ import annotations

import io
import re
import unicodedata
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import (
    AuditLog,
    Classroom,
    CorrectionResult,
    CorrectionRun,
    Evaluation,
    Exercise,
    Participation,
    ScoreAdjustment,
    Subject,
    User,
    utcnow,
)
from .services import participant_scores

HEADER_FILL = PatternFill("solid", fgColor="1D1D1F")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(bold=True, size=14, color="0066CC")
LABEL_FONT = Font(bold=True)
FAIL_FILL = PatternFill("solid", fgColor="FDEEEE")
PASS_FILL = PatternFill("solid", fgColor="E9F8F0")
THIN = Side(style="thin", color="E0E0E0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

RESULT_LABELS = {
    "ok": "Exécutée",
    "compile_error": "Erreur de compilation",
    "runtime_error": "Erreur d'exécution",
    "timeout": "Temps dépassé",
    "no_submission": "Aucune production",
}


def slugify(value: str) -> str:
    text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-zA-Z0-9]+", "-", text).strip("-").lower()[:60] or "evaluation"


def _local(value: datetime | None) -> datetime | None:
    """Excel ne stocke pas de fuseau : on écrit des dates naïves en UTC."""
    return value.replace(tzinfo=None) if value else None


def _header(sheet, row: int, labels: list[str]) -> None:
    for column, label in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=column, value=label)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    sheet.row_dimensions[row].height = 26


def _widths(sheet, widths: list[int]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _duration(seconds: int | None) -> str:
    if seconds is None:
        return "—"
    return f"{seconds // 3600:02d}:{(seconds % 3600) // 60:02d}:{seconds % 60:02d}"


def build_results_workbook(db: Session, evaluation: Evaluation, run: CorrectionRun | None) -> io.BytesIO:
    exercises = list(
        db.scalars(
            select(Exercise).where(Exercise.evaluation_id == evaluation.id).order_by(Exercise.position)
        )
    )
    rows = db.execute(
        select(Participation, User)
        .join(User, User.id == Participation.student_id)
        .where(Participation.evaluation_id == evaluation.id)
        .order_by(User.full_name)
    ).all()
    scores = participant_scores(db, run.id) if run else {}
    results: dict[tuple[int, int], CorrectionResult] = {}
    if run:
        for result in db.scalars(select(CorrectionResult).where(CorrectionResult.run_id == run.id)):
            results[(result.participation_id, result.exercise_id)] = result

    classroom = db.get(Classroom, evaluation.classroom_id) if evaluation.classroom_id else None
    subject = db.get(Subject, evaluation.subject_id) if evaluation.subject_id else None
    teacher = db.get(User, evaluation.teacher_id)

    book = Workbook()
    _synthesis(book.active, evaluation, run, db, classroom, subject, teacher, rows, scores)
    _results(book.create_sheet("Résultats"), evaluation, exercises, rows, scores, results)
    _per_exercise(book.create_sheet("Détail par exercice"), exercises, rows, results)
    _adjustments(book.create_sheet("Ajustements"), db, run, rows)
    _incidents(book.create_sheet("Incidents"), db, evaluation, rows)

    stream = io.BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream


def _synthesis(sheet, evaluation, run, db, classroom, subject, teacher, rows, scores) -> None:
    sheet.title = "Synthèse"
    _widths(sheet, [34, 46])
    sheet["A1"] = "CodEval — résultats d'évaluation"
    sheet["A1"].font = TITLE_FONT

    finals = [scores[p.id]["final"] for p, _ in rows if p.id in scores]
    maxima = [scores[p.id]["max"] for p, _ in rows if p.id in scores]
    barometer = max(maxima) if maxima else evaluation.total_points
    passed = [s for s in finals if barometer and s / barometer >= 0.5]

    entries = [
        ("Évaluation", evaluation.title),
        ("Classe", classroom.name if classroom else "—"),
        ("Matière", subject.name if subject else "—"),
        ("Enseignant", teacher.full_name if teacher else "—"),
        ("Langage", evaluation.language),
        ("Durée (minutes)", evaluation.duration_minutes),
        ("Barème", barometer),
        ("Statut", evaluation.status.value),
        ("Ouverture de la session", _local(evaluation.started_at)),
        ("Clôture de la session", _local(evaluation.closed_at)),
        ("Validation des résultats", _local(evaluation.validated_at)),
        (None, None),
        ("Inscrits", len(rows)),
        ("Ont soumis", sum(1 for p, _ in rows if p.submitted_at)),
        ("Corrigés", len(finals)),
        ("Moyenne", round(sum(finals) / len(finals), 2) if finals else "—"),
        ("Taux de réussite (%)", round(len(passed) / len(finals) * 100, 1) if finals else "—"),
        ("Meilleure note", max(finals) if finals else "—"),
        ("Note la plus basse", min(finals) if finals else "—"),
        ("Notes ajustées manuellement", sum(1 for s in scores.values() if s["adjusted"])),
    ]
    if run:
        entries += [
            (None, None),
            ("Campagne de correction", f"#{run.number}"),
            ("Statut de la campagne", run.status.value),
            ("Lancée le", _local(run.created_at)),
            ("Lancée par", (db.get(User, run.triggered_by).full_name if run.triggered_by else "—")),
            ("Terminée le", _local(run.finished_at)),
            ("Productions traitées", f"{run.processed} / {run.total}"),
            ("Erreurs de compilation", (run.stats or {}).get("compile_errors", 0)),
            ("Dépassements de temps", (run.stats or {}).get("timeouts", 0)),
            ("Sans production", (run.stats or {}).get("no_submission", 0)),
        ]
    else:
        entries += [(None, None), ("Campagne de correction", "aucune correction lancée")]
    entries += [(None, None), ("Export généré le", _local(utcnow()))]

    for index, (label, value) in enumerate(entries, start=3):
        if label is None:
            continue
        sheet.cell(row=index, column=1, value=label).font = LABEL_FONT
        cell = sheet.cell(row=index, column=2, value=value)
        if isinstance(value, datetime):
            cell.number_format = "DD/MM/YYYY HH:MM"


def _results(sheet, evaluation, exercises, rows, scores, results) -> None:
    headers = ["Étudiant", "Matricule", "E-mail"]
    headers += [f"Ex. {e.position} ({e.points:g} pts)" for e in exercises]
    headers += [
        "Note automatique", "Note finale", "Barème", "Ajustée", "Tests réussis", "Tests officiels",
        "Statut", "Temps passé", "Soumission", "Sorties détectées",
    ]
    _header(sheet, 1, headers)
    _widths(sheet, [26, 14, 26] + [13] * len(exercises) + [16, 12, 10, 10, 13, 14, 12, 13, 18, 16])

    for index, (participation, student) in enumerate(rows, start=2):
        entry = scores.get(participation.id)
        line = [student.full_name, student.matricule or "—", student.email]
        passed = total = 0
        for exercise in exercises:
            result = results.get((participation.id, exercise.id))
            line.append(result.auto_score if result else None)
            if result:
                official = [t for t in result.tests if t.get("kind") == "official"]
                total += len(official)
                passed += sum(1 for t in official if t.get("passed"))
        spent = None
        if participation.started_at:
            end = participation.submitted_at or participation.frozen_at
            if end:
                spent = max(0, int((end - participation.started_at).total_seconds()))
        maximum = entry["max"] if entry else evaluation.total_points
        ratio = (entry["final"] / maximum) if entry and maximum else None
        line += [
            entry["auto"] if entry else None,
            entry["final"] if entry else None,
            maximum,
            "oui" if entry and entry["adjusted"] else "non",
            passed,
            total,
            "Non corrigé" if ratio is None else ("Réussi" if ratio >= 0.5 else "Échoué"),
            _duration(spent),
            _local(participation.submitted_at) or "Non soumis",
            participation.incidents or 0,
        ]
        for column, value in enumerate(line, start=1):
            cell = sheet.cell(row=index, column=column, value=value)
            cell.border = BORDER
            if isinstance(value, datetime):
                cell.number_format = "DD/MM/YYYY HH:MM"
        status_cell = sheet.cell(row=index, column=len(headers) - 3)
        if status_cell.value == "Réussi":
            status_cell.fill = PASS_FILL
        elif status_cell.value == "Échoué":
            status_cell.fill = FAIL_FILL

    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(2, len(rows) + 1)}"


def _per_exercise(sheet, exercises, rows, results) -> None:
    _header(sheet, 1, [
        "Étudiant", "Matricule", "Exercice", "Note", "Barème", "État de l'exécution",
        "Tests réussis", "Tests officiels", "Tests échoués", "Message de compilation",
    ])
    _widths(sheet, [26, 14, 40, 10, 10, 24, 14, 15, 40, 60])
    titles = {e.id: f"{e.position}. {e.title}" for e in exercises}

    line = 2
    for participation, student in rows:
        for exercise in exercises:
            result = results.get((participation.id, exercise.id))
            if result is None:
                continue
            official = [t for t in result.tests if t.get("kind") == "official"]
            failed = [t["name"] for t in result.tests if not t.get("passed")]
            values = [
                student.full_name,
                student.matricule or "—",
                titles.get(exercise.id, ""),
                result.auto_score,
                result.max_score,
                RESULT_LABELS.get(result.status.value, result.status.value),
                sum(1 for t in official if t.get("passed")),
                len(official),
                ", ".join(failed) or "—",
                (result.compile_log or "")[:500],
            ]
            for column, value in enumerate(values, start=1):
                cell = sheet.cell(row=line, column=column, value=value)
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=column >= 9)
            if failed:
                sheet.cell(row=line, column=6).fill = FAIL_FILL
            line += 1
    sheet.freeze_panes = "A2"


def _adjustments(sheet, db, run, rows) -> None:
    _header(sheet, 1, ["Date", "Étudiant", "Enseignant", "Note précédente", "Nouvelle note", "Motif"])
    _widths(sheet, [20, 26, 24, 16, 15, 60])
    if run is None:
        return
    names = {p.id: s.full_name for p, s in rows}
    line = 2
    for adjustment in db.scalars(
        select(ScoreAdjustment)
        .where(ScoreAdjustment.run_id == run.id)
        .order_by(ScoreAdjustment.created_at)
    ):
        author = db.get(User, adjustment.teacher_id)
        values = [
            _local(adjustment.created_at),
            names.get(adjustment.participation_id, "—"),
            author.full_name if author else "—",
            adjustment.previous_score,
            adjustment.new_score,
            adjustment.reason,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=line, column=column, value=value)
            cell.border = BORDER
            if isinstance(value, datetime):
                cell.number_format = "DD/MM/YYYY HH:MM"
        line += 1
    sheet.freeze_panes = "A2"


def _incidents(sheet, db, evaluation, rows) -> None:
    _header(sheet, 1, ["Date", "Étudiant", "Type d'incident"])
    _widths(sheet, [20, 26, 26])
    names = {p.id: s.full_name for p, s in rows}
    labels = {
        "fullscreen_exit": "Sortie du plein écran",
        "tab_hidden": "Changement d'onglet",
        "window_blur": "Perte de focus",
        "paste_blocked": "Collage bloqué",
        "copy_blocked": "Copie bloquée",
        "shortcut_blocked": "Raccourci bloqué",
        "reload_attempt": "Tentative de rechargement",
    }
    line = 2
    for entry in db.scalars(
        select(AuditLog)
        .where(
            AuditLog.organization_id == evaluation.organization_id,
            AuditLog.action == "integrity.incident",
            AuditLog.target_type == "participation",
            AuditLog.target_id.in_(names.keys() or [0]),
        )
        .order_by(AuditLog.created_at)
    ):
        kind = (entry.meta or {}).get("type", "")
        for column, value in enumerate(
            [_local(entry.created_at), names.get(entry.target_id, "—"), labels.get(kind, kind)],
            start=1,
        ):
            cell = sheet.cell(row=line, column=column, value=value)
            cell.border = BORDER
            if isinstance(value, datetime):
                cell.number_format = "DD/MM/YYYY HH:MM"
        line += 1
    sheet.freeze_panes = "A2"
